"""Building the school from student lists, the bell periods from the bell system, teachers from the venues /
tickets exports, and the news ticker."""
import io
import re

import pytest
from openpyxl import Workbook

from extensions import db
from models import Stage, Grade, Section, Student, Period, TimetableEntry, Employee, Tenant
from tenancy import tenant_context
from tests.test_school import school, _client  # noqa: F401  (fixture)


@pytest.mark.parametrize('text, expected', [
    ('4A', (4, 'A')), ('4CPA', (4, 'CP A')), ('4CPB NEW ', (4, 'CP B')), ('Grade 6-C', (6, 'C')),
    ('Grade7A', (7, 'A')), ('Five CP-A', (5, 'CP A')), ('Seven CS-A', (7, 'CS A')),
    ('Nine IGCSE A', (9, 'IGCSE A')), ('Twelve A2 - Scientific', (12, 'A2 Scientific')),
    ('4CP- A', (4, 'CP A')), ('الصف الرابع / أ', (4, 'أ')), ('الصف الحادي عشر ب', (11, 'ب')),
    ('Eleven AS Literary', (11, 'AS Literary')),
])
def test_class_codes(text, expected):
    from services.importing import parse_class
    assert parse_class(text) == expected


def _xlsx(sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _preview(c, data, name='list.xlsx'):
    r = c.post('/students/import', data={'file': (data, name)}, content_type='multipart/form-data')
    assert r.status_code == 200, r.status_code
    html = r.get_data(as_text=True)
    token = re.search(r'/students/import/([0-9a-f]{32})/apply', html).group(1)
    return html, token


def _school_sheets():
    """Like the school's grade-4 workbook: a title line, the header on line 3, one sheet per section with the
    class code in the «Grade» column, and copy sheets for the notebooks list."""
    def sheet(code, names):
        return [['Last Update :28/8/2024'], ['Num.', 'Std. Name', 'Grade']] + \
            [[i, n, code] for i, n in enumerate(names, 1)]
    a = ['طالب أول', 'طالبة ثانية', 'طالب ثالث']
    b = ['طالب رابع', 'طالب خامس']
    cpa = ['طالبة سادسة', 'طالب سابع']
    return [('4-A', sheet('4A', a)), ('4- B', sheet('4B', b)), ('4CP- A', sheet('4CPA', cpa)),
            ('4-A للدفاتر', sheet('4A', a[:2])), ('4- B للدفاتر', sheet('4B', b))]


def test_students_import_builds_the_school_and_is_idempotent(app, school):
    with tenant_context(school['a']):
        st = Stage(name='المرحلة الأساسية 4-6', sort_order=5)
        db.session.add(st)
        db.session.commit()
        stage_id = st.id
    c = _client(app, 'admin')
    html, token = _preview(c, _xlsx(_school_sheets()))
    with tenant_context(school['a']):
        assert Grade.query.count() == 2                              # nothing written by the preview
    assert 'نسخة — تُركت' in html                                     # the notebook copies are left out
    c.post(f'/students/import/{token}/apply')
    with tenant_context(school['a']):
        g4 = Grade.query.filter_by(name='الصف الرابع').one()
        assert g4.stage_id == stage_id and g4.name_en == 'Grade 4'  # the stage whose name covers grade 4
        assert sorted(s.name for s in g4.sections) == ['A', 'B', 'CP A']
        assert Student.query.join(Section).filter(Section.grade_id == g4.id).count() == 7
    # the same file again: nothing new
    html, token = _preview(c, _xlsx(_school_sheets()))
    assert re.search(r'<div class="n">0</div><div class="l">طلبة جدد', html)
    c.post(f'/students/import/{token}/apply')
    with tenant_context(school['a']):
        assert Student.query.join(Section).join(Grade).filter(Grade.name == 'الصف الرابع').count() == 7


def test_students_import_new_stage_section_names_and_moves(app, school):
    rows = [['المرحلة', 'الصف', 'الشعبة', 'الشعبة بالإنجليزية', 'الاسم', 'مربي الصف'],
            ['المرحلة الثانوية', 'الصف الثاني عشر', 'الحقل الصحي', 'Academic - Health', 'Test Student One', '(Empty)'],
            ['المرحلة الثانوية', 'Twelve A2 - Scientific', '', '', 'Test Student Two', 'Ms. Huda'],
            ['', 'الصف 2', 'أ', '', 'طالب 1-1', '']]                 # an existing student, other section
    c = _client(app, 'admin')
    buf = io.BytesIO()
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    wb.save(buf)
    buf.seek(0)
    r = c.post('/students/import', data={'file': (buf, 'x.xlsx'), 'update': '1'},
               content_type='multipart/form-data')
    token = re.search(r'/students/import/([0-9a-f]{32})/apply', r.get_data(as_text=True)).group(1)
    c.post(f'/students/import/{token}/apply')
    with tenant_context(school['a']):
        st = Stage.query.filter_by(name='المرحلة الثانوية').one()
        g12 = Grade.query.filter_by(stage_id=st.id).one()
        health = Section.query.filter_by(grade_id=g12.id, name='الحقل الصحي').one()
        assert health.name_en == 'Academic - Health' and health.homeroom_teacher is None
        a2 = Section.query.filter_by(grade_id=g12.id, name='A2 Scientific').one()
        assert a2.homeroom_teacher == 'Ms. Huda'
        hala = Student.query.filter_by(full_name='Test Student Two').one()
        assert hala.full_name_en == 'Test Student Two' and hala.section_id == a2.id
        assert db.session.get(Student, school['st11']).section_id == school['sec2']    # moved


def test_secretary_cannot_create_structure_by_import(app, school):
    c = _client(app, 'secy')
    html, token = _preview(c, _xlsx([('s', [['الاسم', 'الصف'], ['طالب جديد', '9B']])]))
    assert 'disabled' in html
    assert c.post(f'/students/import/{token}/apply').status_code == 403
    # another user cannot use the preview of someone else
    other = _client(app, 'admin')
    r = other.post(f'/students/import/{token}/apply')
    assert r.status_code == 302 and '/students/import' in r.headers['Location']


def _bell_rows():
    head = ['المرحلة', 'المرحلة بالإنجليزية', 'الصف', 'الصف بالإنجليزية', 'رقم اليوم', 'اليوم', 'الترتيب', 'النوع',
            'الاسم', 'الاسم بالإنجليزية', 'تبدأ', 'تنتهي']
    rows = [head]
    p56 = [('حصة', 'الأولى', '08:00', '08:45'), ('حصة', 'الثانية', '08:45', '09:30'),
           ('استراحة', 'الفرصة', '09:30', '09:50'), ('حصة', 'الثالثة', '09:50', '10:35')]
    p4 = [('حصة', 'الأولى', '08:00', '08:40'), ('استراحة', 'الفرصة', '08:40', '09:00'),
          ('حصة', 'الثانية', '09:00', '09:45')]
    for grade, plist in (('الصف 1', p56), ('الصف 3', p56), ('الصف 4', p4), ('الصف التاسع', p56)):
        for wd in (6, 0):
            for i, (k, n, s, e) in enumerate(plist, 1):
                rows.append(['المرحلة 1', '', grade, '', wd, '', i, k, n, '', s, e])
    return rows


def test_periods_from_the_bell_system(app, school):
    with tenant_context(school['a']):
        for name in ('الصف 3', 'الصف 4'):
            db.session.add(Grade(stage_id=school['stage1'], name=name))
        db.session.commit()
        g1 = db.session.get(Grade, school['grade1'])
        old = Period(stage_id=g1.stage_id, weekday=6, sort_order=1, label='قديمة', start_time='07:30',
                     end_time='08:00')
        db.session.add(old)
        db.session.flush()
        db.session.add(TimetableEntry(section_id=school['sec1'], period_id=old.id, subject='رياضيات'))
        db.session.commit()
        old_id = old.id
    c = _client(app, 'admin')
    r = c.post('/structure/periods/import', data={'file': (_xlsx([('الحصص', _bell_rows())]), 'bell.xlsx')},
               content_type='multipart/form-data')
    html = r.get_data(as_text=True)
    assert 'الصف التاسع' in html                                  # not in SAS: left out, said so
    token = re.search(r'/structure/periods/import/([0-9a-f]{32})/apply', html).group(1)
    c.post(f'/structure/periods/import/{token}/apply', data={'school_days': '1'})
    from services.periods import day_periods
    with tenant_context(school['a']):
        g1 = db.session.get(Grade, school['grade1'])
        g4 = Grade.query.filter_by(name='الصف 4').one()
        assert [p.start_time for p in day_periods(g1, 6)] == ['08:00', '08:45', '09:30', '09:50']
        assert [p.kind for p in day_periods(g4, 6)] == ['period', 'break', 'period']   # grade 4: its own day
        assert Period.query.filter_by(grade_id=g1.id).count() == 0                     # same as the stage
        kept = db.session.get(Period, old_id)                                          # updated in place
        assert kept is not None and kept.start_time == '08:00'
        assert TimetableEntry.query.filter_by(period_id=old_id).count() == 1
        assert db.session.get(Stage, school['stage1']).school_days == '0,6'
    assert c.get(f"/timetable/section/{school['sec1']}").status_code == 200
    # importing again changes nothing
    r = c.post('/structure/periods/import', data={'file': (_xlsx([('الحصص', _bell_rows())]), 'bell.xlsx')},
               content_type='multipart/form-data')
    token = re.search(r'/structure/periods/import/([0-9a-f]{32})/apply', r.get_data(as_text=True)).group(1)
    r = c.post(f'/structure/periods/import/{token}/apply', follow_redirects=True)
    assert 'أُضيف 0، وعُدّل 0، وحُذف 0' in r.get_data(as_text=True)


def test_teachers_from_venues_and_tickets_merge(app, school):
    venues = '﻿الاسم,البريد الإلكتروني,الهاتف,المسمى الوظيفي,المرحلة,الصف,الشعبة,الحالة,المصدر\n' \
             'معلمة تجربة,,0791111111,معلم,المرحلة الأساسية 4-6,,,فعّال,نظام القاعات\n' \
             'معلم آخر,mona@x.test,,معلم,,,,فعّال,نظام القاعات\n'
    tickets = '﻿الاسم,البريد الإلكتروني,الهاتف,المسمى الوظيفي,التخصص,الحالة,المصدر\n' \
              'معلمة  تجربة,teacher@x.test,,معلم,العلوم,فعّال,نظام التذاكر\n'
    c = _client(app, 'admin')
    for text in (venues, tickets):
        c.post('/employees/import', data={'file': (io.BytesIO(text.encode('utf-8')), 'teachers-for-SAS.csv')},
               content_type='multipart/form-data')
    with tenant_context(school['a']):
        alaa = Employee.query.filter(Employee.full_name.like('معلمة%')).all()
        assert len(alaa) == 1                                         # one person, two sources
        assert alaa[0].email == 'teacher@x.test' and alaa[0].phone == '0791111111'
        assert alaa[0].job_title == 'معلم'                            # not overwritten
        assert Employee.query.count() == 2


def test_ticker_settings_and_page(app, school):
    c = _client(app, 'admin')
    assert c.get('/admin/ticker').status_code == 200
    r = c.post('/admin/ticker', data={'enabled': '1', 'feeds_ar': 'أهلاً <script>alert(1)</script>\nالثانية',
                                      'feeds_en': 'Welcome', 'fg': '#112233', 'bg': '#ffffff', 'opacity': '80',
                                      'size': '99', 'speed': '40', 'direction': 'auto', 'separator': 'diamond',
                                      'logo': 'school', 'on_login': '1', 'font': 'evil;}'})
    assert r.status_code == 302
    with tenant_context(school['a']):
        import json
        cfg = json.loads(db.session.get(Tenant, school['a']).ticker_json)
        assert cfg['size'] == 32 and cfg['font'] == 'inherit' and cfg['opacity'] == 80
    html = c.get('/home').get_data(as_text=True)
    assert 'data-ticker' in html and 'data-move="right"' in html and '◆' in html
    assert '<script>alert(1)' not in html and '&lt;script&gt;' in html
    c.get('/lang/en')
    html = c.get('/home').get_data(as_text=True)
    assert 'Welcome' in html and 'data-move="left"' in html
    # an invalid colour is refused; turning it off removes it from the pages
    r = c.post('/admin/ticker', data={'enabled': '1', 'feeds_ar': 'x', 'fg': 'red'})
    assert r.status_code == 400
    c.post('/admin/ticker', data={'feeds_ar': 'x'})
    assert 'data-ticker' not in c.get('/home').get_data(as_text=True)
