"""Excel (.xlsx) download of a list: right-to-left sheet, bold frozen header, sensible widths."""
import io

from flask import Response


def xlsx_response(filename, headers, rows, sheet='البيانات'):
    """Headers and the sheet name in the page language (left-to-right sheet in English)."""
    from openpyxl import Workbook
    from utils.i18n import _, is_en
    headers = [_(h) for h in headers]
    sheet = _(sheet)
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    ws.sheet_view.rightToLeft = not is_en()
    ws.append(list(headers))
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1B3358')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append(['' if v is None else (v if isinstance(v, (int, float)) else str(v)) for v in r])
    ws.freeze_panes = 'A2'
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or '')) for r in rows[:300]]) + 3
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width, 8), 45)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition': f'attachment; filename={filename}'})
